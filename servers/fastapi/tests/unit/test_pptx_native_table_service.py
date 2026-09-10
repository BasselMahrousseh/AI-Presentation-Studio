import asyncio
import io
import uuid

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from services import pptx_native_table_service as svc

_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def _make_presentation_with_picture(left_px, top_px, width_px, height_px):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    emu_x = prs.slide_width / 1280
    emu_y = prs.slide_height / 720

    img = Image.new("RGB", (10, 10), color="red")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    picture = slide.shapes.add_picture(
        buf,
        Emu(int(left_px * emu_x)),
        Emu(int(top_px * emu_y)),
        Emu(int(width_px * emu_x)),
        Emu(int(height_px * emu_y)),
    )
    return prs, slide, picture, emu_x, emu_y


def _simple_table(row_count=2, col_count=2, **overrides):
    cells = overrides.pop("cells", None)
    if cells is None:
        cells = [
            {"rowIndex": r, "colIndex": c, "rowSpan": 1, "colSpan": 1, "text": f"r{r}c{c}"}
            for r in range(row_count)
            for c in range(col_count)
        ]
    table = {"rowCount": row_count, "colCount": col_count, "cells": cells}
    table.update(overrides)
    return table


# ---------------------------------------------------------------------------
# Font weight -> bold heuristic
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "value,expected",
    [
        (700, True),
        (600, True),
        (599, False),
        ("700", True),
        ("bold", True),
        ("normal", False),
        ("garbage", False),
        (None, False),
        (True, False),
    ],
)
def test_font_weight_is_bold(value, expected):
    assert svc._font_weight_is_bold(value) is expected


# ---------------------------------------------------------------------------
# _build_table_grid validation
# ---------------------------------------------------------------------------


def test_build_table_grid_happy_path_normalizes_and_sorts():
    grid = svc._build_table_grid(_simple_table(2, 2))
    assert grid is not None
    assert grid["rowCount"] == 2
    assert grid["colCount"] == 2
    assert len(grid["cells"]) == 4


@pytest.mark.parametrize(
    "overrides",
    [
        {"rowCount": 0},
        {"rowCount": -1},
        {"rowCount": "2"},
        {"rowCount": True},
        {"colCount": 0},
    ],
)
def test_build_table_grid_rejects_bad_row_or_col_count(overrides):
    table = _simple_table(2, 2)
    table.update(overrides)
    assert svc._build_table_grid(table) is None


def test_build_table_grid_rejects_missing_or_empty_cells():
    assert svc._build_table_grid({"rowCount": 2, "colCount": 2}) is None
    assert svc._build_table_grid({"rowCount": 2, "colCount": 2, "cells": []}) is None
    assert svc._build_table_grid({"rowCount": 2, "colCount": 2, "cells": "nope"}) is None


def test_build_table_grid_rejects_non_dict_cell():
    table = _simple_table(1, 1, cells=["not-a-dict"])
    assert svc._build_table_grid(table) is None


@pytest.mark.parametrize(
    "cell_overrides",
    [
        {"rowSpan": 0},
        {"colSpan": 0},
        {"rowSpan": -1},
        {"rowIndex": -1},
        {"colIndex": -1},
        {"rowIndex": "0"},
        {"rowSpan": True},
    ],
)
def test_build_table_grid_rejects_bad_span_or_index_values(cell_overrides):
    cell = {"rowIndex": 0, "colIndex": 0, "rowSpan": 1, "colSpan": 1, "text": "x"}
    cell.update(cell_overrides)
    table = {"rowCount": 1, "colCount": 1, "cells": [cell]}
    assert svc._build_table_grid(table) is None


def test_build_table_grid_rejects_span_extending_out_of_bounds():
    cell = {"rowIndex": 0, "colIndex": 0, "rowSpan": 3, "colSpan": 1}
    table = {"rowCount": 2, "colCount": 2, "cells": [cell]}
    assert svc._build_table_grid(table) is None


def test_build_table_grid_rejects_overlapping_anchors():
    cells = [
        {"rowIndex": 0, "colIndex": 0, "rowSpan": 1, "colSpan": 1, "text": "a"},
        {"rowIndex": 0, "colIndex": 0, "rowSpan": 1, "colSpan": 1, "text": "b"},
    ]
    table = {"rowCount": 1, "colCount": 2, "cells": cells}
    assert svc._build_table_grid(table) is None


def test_build_table_grid_rejects_over_the_cell_cap():
    # 501 cells, exceeding _MAX_TABLE_CELLS = 500.
    table = {"rowCount": 1, "colCount": svc._MAX_TABLE_CELLS + 1, "cells": []}
    assert svc._build_table_grid(table) is None


def test_build_table_grid_accepts_exactly_the_cell_cap():
    row_count, col_count = 1, svc._MAX_TABLE_CELLS
    cells = [
        {"rowIndex": 0, "colIndex": c, "rowSpan": 1, "colSpan": 1, "text": str(c)}
        for c in range(col_count)
    ]
    table = {"rowCount": row_count, "colCount": col_count, "cells": cells}
    assert svc._build_table_grid(table) is not None


# ---------------------------------------------------------------------------
# Header-row detection
# ---------------------------------------------------------------------------


def test_looks_like_header_row_true_when_flagged():
    cells = [{"rowIndex": 0, "colIndex": 0, "isHeaderCell": True}]
    assert svc._looks_like_header_row(cells, col_count=1) is True


def test_looks_like_header_row_true_on_bold_majority():
    cells = [
        {"rowIndex": 0, "colIndex": 0, "fontWeight": 700},
        {"rowIndex": 0, "colIndex": 1, "fontWeight": 700},
        {"rowIndex": 0, "colIndex": 2, "fontWeight": 400},
    ]
    assert svc._looks_like_header_row(cells, col_count=3) is True


def test_looks_like_header_row_false_when_plain():
    cells = [
        {"rowIndex": 0, "colIndex": 0, "fontWeight": 400},
        {"rowIndex": 0, "colIndex": 1, "fontWeight": 400},
    ]
    assert svc._looks_like_header_row(cells, col_count=2) is False


def test_looks_like_header_row_false_with_no_row_zero_cells():
    cells = [{"rowIndex": 1, "colIndex": 0, "isHeaderCell": True}]
    assert svc._looks_like_header_row(cells, col_count=1) is False


# ---------------------------------------------------------------------------
# End-to-end: build a real native table from a synthetic capture
# ---------------------------------------------------------------------------


def test_try_upgrade_one_table_replaces_matching_picture_with_native_table():
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(100, 100, 400, 300)
    captured_table = {
        "slideOrderIndex": 0,
        "boundingBox": {"left": 100, "top": 100, "width": 400, "height": 300},
        "rowCount": 2,
        "colCount": 2,
        "cells": [
            {"rowIndex": 0, "colIndex": 0, "rowSpan": 1, "colSpan": 1, "text": "A"},
            {"rowIndex": 0, "colIndex": 1, "rowSpan": 1, "colSpan": 1, "text": "B"},
            {"rowIndex": 1, "colIndex": 0, "rowSpan": 1, "colSpan": 1, "text": "C"},
            {"rowIndex": 1, "colIndex": 1, "rowSpan": 1, "colSpan": 1, "text": "D"},
        ],
    }
    claimed: set = set()
    upgraded = svc._try_upgrade_one_table(slide, captured_table, emu_x, emu_y, claimed)
    assert upgraded is True

    shapes = list(slide.shapes)
    assert len(shapes) == 1
    assert shapes[0].shape_type == MSO_SHAPE_TYPE.TABLE
    table = shapes[0].table
    assert table.cell(0, 0).text == "A"
    assert table.cell(1, 1).text == "D"


def test_try_upgrade_one_table_unsupported_grid_leaves_picture_untouched():
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(100, 100, 400, 300)
    captured_table = {
        "boundingBox": {"left": 100, "top": 100, "width": 400, "height": 300},
        "rowCount": 0,  # invalid
        "colCount": 2,
        "cells": [],
    }
    claimed: set = set()
    upgraded = svc._try_upgrade_one_table(slide, captured_table, emu_x, emu_y, claimed)
    assert upgraded is False
    shapes = list(slide.shapes)
    assert len(shapes) == 1
    assert shapes[0].shape_type == MSO_SHAPE_TYPE.PICTURE


def test_try_upgrade_one_table_no_confident_match_leaves_picture_untouched():
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(100, 100, 400, 300)
    captured_table = _simple_table(2, 2)
    captured_table["boundingBox"] = {"left": 900, "top": 600, "width": 50, "height": 50}
    claimed: set = set()
    upgraded = svc._try_upgrade_one_table(slide, captured_table, emu_x, emu_y, claimed)
    assert upgraded is False
    shapes = list(slide.shapes)
    assert shapes[0].shape_type == MSO_SHAPE_TYPE.PICTURE


def test_try_upgrade_one_table_missing_bounding_box_leaves_picture_untouched():
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(100, 100, 400, 300)
    captured_table = _simple_table(2, 2)
    claimed: set = set()
    upgraded = svc._try_upgrade_one_table(slide, captured_table, emu_x, emu_y, claimed)
    assert upgraded is False


# ---------------------------------------------------------------------------
# Failure path: add_table() raising must restore the original picture
# ---------------------------------------------------------------------------


def test_failed_add_table_leaves_the_original_picture_in_place(monkeypatch):
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(100, 100, 400, 300)
    original_element = picture._element
    original_index = list(slide.shapes._spTree).index(original_element)

    def _boom(*args, **kwargs):
        raise RuntimeError("malformed table data")

    monkeypatch.setattr(slide.shapes, "add_table", _boom)

    captured_table = svc._build_table_grid(_simple_table(2, 2))
    upgraded = svc._replace_picture_with_native_table(slide, picture, captured_table)

    assert upgraded is False
    shapes = list(slide.shapes)
    assert len(shapes) == 1
    assert shapes[0].shape_type == MSO_SHAPE_TYPE.PICTURE
    assert shapes[0]._element is original_element
    sp_tree = slide.shapes._spTree
    assert list(sp_tree).index(original_element) == original_index


def test_failed_content_application_after_add_table_succeeds_still_restores_picture(
    monkeypatch,
):
    """A structural failure (not just add_table() itself raising) after the
    table shape was already created and repositioned must still roll back
    completely - both the picture restored AND the half-built table shape
    removed, not left behind as an extra broken shape."""
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(100, 100, 400, 300)
    original_element = picture._element

    def _boom(*args, **kwargs):
        raise RuntimeError("bad content")

    monkeypatch.setattr(svc, "_apply_table_content", _boom)

    captured_table = svc._build_table_grid(_simple_table(2, 2))
    upgraded = svc._replace_picture_with_native_table(slide, picture, captured_table)

    assert upgraded is False
    shapes = list(slide.shapes)
    assert len(shapes) == 1
    assert shapes[0].shape_type == MSO_SHAPE_TYPE.PICTURE
    assert shapes[0]._element is original_element


def test_styling_failure_still_leaves_a_structurally_correct_table(monkeypatch):
    """Styling is best-effort/cosmetic - unlike content/merge, a failure here
    must NOT roll back the already-successfully-built table."""
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(100, 100, 400, 300)

    def _boom(*args, **kwargs):
        raise RuntimeError("styling exploded")

    monkeypatch.setattr(svc, "_apply_table_styling", _boom)

    captured_table = svc._build_table_grid(_simple_table(2, 2))
    upgraded = svc._replace_picture_with_native_table(slide, picture, captured_table)

    assert upgraded is True
    shapes = list(slide.shapes)
    assert shapes[0].shape_type == MSO_SHAPE_TYPE.TABLE


# ---------------------------------------------------------------------------
# Preserves z-order (position in the shape tree)
# ---------------------------------------------------------------------------


def test_try_upgrade_one_table_preserves_z_order():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    emu_x = prs.slide_width / 1280
    emu_y = prs.slide_height / 720

    img = Image.new("RGB", (10, 10), color="blue")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)

    before = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(100), Emu(100))
    before.text_frame.text = "before"
    buf.seek(0)
    picture = slide.shapes.add_picture(
        buf, Emu(int(100 * emu_x)), Emu(int(100 * emu_y)), Emu(int(400 * emu_x)), Emu(int(300 * emu_y))
    )
    after = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(100), Emu(100))
    after.text_frame.text = "after"

    captured_table = _simple_table(2, 2)
    captured_table["boundingBox"] = {"left": 100, "top": 100, "width": 400, "height": 300}
    claimed: set = set()
    upgraded = svc._try_upgrade_one_table(slide, captured_table, emu_x, emu_y, claimed)
    assert upgraded is True

    shapes = list(slide.shapes)
    assert len(shapes) == 3
    assert shapes[0].text_frame.text == "before"
    assert shapes[1].shape_type == MSO_SHAPE_TYPE.TABLE
    assert shapes[2].text_frame.text == "after"


# ---------------------------------------------------------------------------
# Leftover per-cell fragment cleanup - a real bug found via a live user
# report: the export converter doesn't always flatten a <table> into one
# whole-table picture. For a table with per-cell background colors/borders,
# it also emits a separate PICTURE (that cell's background/border) plus a
# separate TEXT_BOX (that cell's real text) for every cell, layered on top
# of/alongside the one whole-table backdrop picture this pipeline matches
# and swaps. Left unhandled, those per-cell fragments stayed behind, visibly
# duplicating every cell's text in the real exported file (confirmed via a
# real user screenshot and reproduced by inspecting the actual broken file's
# full shape list - python-pptx's own .text on the TABLE object was always
# correct, which is exactly why this went unnoticed until someone opened the
# real file and looked at more than just the table shape itself).
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "outer,inner,expected",
    [
        ((0, 0, 100, 100), (10, 10, 20, 20), 1.0),  # fully inside
        ((0, 0, 100, 100), (50, 50, 100, 100), 0.25),  # half in, half out
        ((0, 0, 100, 100), (200, 200, 10, 10), 0.0),  # no overlap
        ((0, 0, 100, 100), (0, 0, 0, 0), 0.0),  # degenerate inner
    ],
)
def test_containment_ratio(outer, inner, expected):
    assert svc._containment_ratio(outer, inner) == pytest.approx(expected)


def _add_picture(slide, left, top, width, height, color="blue"):
    img = Image.new("RGB", (10, 10), color=color)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return slide.shapes.add_picture(buf, Emu(left), Emu(top), Emu(width), Emu(height))


def test_remove_leftover_table_fragments_removes_contained_pictures_and_textboxes():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table_rect = (0, 0, 1000, 1000)
    graphic_frame = slide.shapes.add_table(2, 2, Emu(0), Emu(0), Emu(1000), Emu(1000))

    # Two per-cell fragments fully inside the table's rect.
    frag_pic = _add_picture(slide, 0, 0, 400, 400)
    frag_box = slide.shapes.add_textbox(Emu(500), Emu(0), Emu(400), Emu(400))
    # A shape mostly outside the table's rect must survive.
    outside_pic = _add_picture(slide, 900, 900, 500, 500)

    removed = svc._remove_leftover_table_fragments(
        slide, table_rect, graphic_frame._element
    )

    assert removed == 2
    remaining_elements = {id(s._element) for s in slide.shapes}
    assert id(graphic_frame._element) in remaining_elements
    assert id(outside_pic._element) in remaining_elements
    assert id(frag_pic._element) not in remaining_elements
    assert id(frag_box._element) not in remaining_elements


def test_replace_picture_with_native_table_cleans_up_leftover_fragments_end_to_end():
    """Reproduces the real bug directly: a picture representing a table's
    whole-backdrop screenshot, PLUS several per-cell fragment pictures the
    converter also emitted alongside it (fully contained within the same
    rect) - matching the real shape layout found in the actual broken
    exported file. After conversion, none of those fragments should remain,
    and unrelated content elsewhere on the slide must be untouched."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    backdrop = _add_picture(slide, 0, 0, 1000, 1000, color="white")
    # Per-cell fragments the converter also emitted for this same table,
    # each fully inside the backdrop's rect - the exact pattern observed in
    # the real broken file (a PICTURE + TEXT_BOX pair per cell).
    for row in range(2):
        for col in range(2):
            left, top = col * 500, row * 500
            _add_picture(slide, left, top, 500, 500, color="red")
            slide.shapes.add_textbox(Emu(left), Emu(top), Emu(500), Emu(500))
    unrelated = slide.shapes.add_textbox(Emu(2000), Emu(2000), Emu(100), Emu(100))
    unrelated.text_frame.text = "unrelated content"

    captured_table = _simple_table(2, 2)
    captured_table["boundingBox"] = {"left": 0, "top": 0, "width": 1000, "height": 1000}
    # 1000 EMU-equivalent px maps 1:1 here since emu_per_px=1 for this test.
    upgraded = svc._replace_picture_with_native_table(slide, backdrop, captured_table)
    assert upgraded is True

    shapes = list(slide.shapes)
    assert sum(1 for s in shapes if s.shape_type == MSO_SHAPE_TYPE.TABLE) == 1
    assert sum(1 for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE) == 0
    # Only the table and the untouched unrelated textbox should remain.
    text_boxes = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]
    assert len(text_boxes) == 1
    assert text_boxes[0].text_frame.text == "unrelated content"


def test_rollback_path_never_removes_fragments():
    """If the table conversion itself fails and the original picture is
    restored, the picture-based fallback rendering - including whatever
    fragments compose it - must be left completely untouched. Fragment
    cleanup only ever runs on the success path."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    backdrop = _add_picture(slide, 0, 0, 1000, 1000, color="white")
    fragment = _add_picture(slide, 0, 0, 500, 500, color="red")

    def _boom(*args, **kwargs):
        raise RuntimeError("bad content")

    with_patch = pytest.MonkeyPatch()
    with_patch.setattr(svc, "_apply_table_content", _boom)
    try:
        captured_table = svc._build_table_grid(_simple_table(2, 2))
        upgraded = svc._replace_picture_with_native_table(slide, backdrop, captured_table)
    finally:
        with_patch.undo()

    assert upgraded is False
    shapes = list(slide.shapes)
    assert len(shapes) == 2
    assert any(s._element is backdrop._element for s in shapes)
    assert any(s._element is fragment._element for s in shapes)


# ---------------------------------------------------------------------------
# Merging - emitted OOXML, not just python-pptx properties (this codebase's
# own stated discipline given the dLblPos corruption history: python-pptx
# round-tripping a file successfully has been proven insufficient before).
# ---------------------------------------------------------------------------


def test_merge_emits_gridspan_and_rowspan_hmerge_vmerge_and_survives_reload(tmp_path):
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(100, 100, 400, 300)
    captured_table = {
        "boundingBox": {"left": 100, "top": 100, "width": 400, "height": 300},
        "rowCount": 2,
        "colCount": 2,
        "cells": [
            {"rowIndex": 0, "colIndex": 0, "rowSpan": 1, "colSpan": 2, "text": "Wide header"},
            {"rowIndex": 1, "colIndex": 0, "rowSpan": 1, "colSpan": 1, "text": "A"},
            {"rowIndex": 1, "colIndex": 1, "rowSpan": 1, "colSpan": 1, "text": "B"},
        ],
    }
    claimed: set = set()
    upgraded = svc._try_upgrade_one_table(slide, captured_table, emu_x, emu_y, claimed)
    assert upgraded is True

    pptx_path = tmp_path / "merge.pptx"
    prs.save(str(pptx_path))

    reopened = Presentation(str(pptx_path))
    tbl_el = reopened.slides[0].shapes[0].table._tbl
    tcs = tbl_el.findall(".//a:tr", _NS)[0].findall("a:tc", _NS)
    assert tcs[0].get("gridSpan") == "2"
    assert tcs[1].get("hMerge") == "1"

    reopened_table = reopened.slides[0].shapes[0].table
    assert reopened_table.cell(0, 0).text == "Wide header"
    assert reopened_table.cell(1, 0).text == "A"
    assert reopened_table.cell(1, 1).text == "B"


def test_rowspan_and_colspan_combined_merge_survives_reload(tmp_path):
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(100, 100, 400, 300)
    captured_table = {
        "boundingBox": {"left": 100, "top": 100, "width": 400, "height": 300},
        "rowCount": 3,
        "colCount": 3,
        "cells": [
            {"rowIndex": 0, "colIndex": 0, "rowSpan": 2, "colSpan": 2, "text": "Block"},
            {"rowIndex": 0, "colIndex": 2, "rowSpan": 1, "colSpan": 1, "text": "H3"},
            {"rowIndex": 1, "colIndex": 2, "rowSpan": 1, "colSpan": 1, "text": "C1"},
            {"rowIndex": 2, "colIndex": 0, "rowSpan": 1, "colSpan": 1, "text": "R2C0"},
            {"rowIndex": 2, "colIndex": 1, "rowSpan": 1, "colSpan": 1, "text": "R2C1"},
            {"rowIndex": 2, "colIndex": 2, "rowSpan": 1, "colSpan": 1, "text": "R2C2"},
        ],
    }
    claimed: set = set()
    upgraded = svc._try_upgrade_one_table(slide, captured_table, emu_x, emu_y, claimed)
    assert upgraded is True

    pptx_path = tmp_path / "combo_merge.pptx"
    prs.save(str(pptx_path))

    reopened = Presentation(str(pptx_path))
    trs = reopened.slides[0].shapes[0].table._tbl.findall("a:tr", _NS)
    row0_tcs = trs[0].findall("a:tc", _NS)
    row1_tcs = trs[1].findall("a:tc", _NS)

    assert row0_tcs[0].get("gridSpan") == "2"
    assert row0_tcs[0].get("rowSpan") == "2"
    assert row0_tcs[1].get("hMerge") == "1"
    assert row1_tcs[0].get("vMerge") == "1"
    assert row1_tcs[1].get("hMerge") == "1"
    assert row1_tcs[1].get("vMerge") == "1"


def test_invalid_merge_is_skipped_without_raising():
    """A merge whose footprint collides with an already-merged region raises
    ValueError from python-pptx's own .merge() (confirmed directly against
    the real API: "range contains one or more merged cells") - that one bad
    span must be skipped, not sink the whole table. Two cells with distinct
    anchors but overlapping footprints (only the anchor is checked for
    collisions, not the full merged footprint) is what produces this for
    real, without needing to fake anything."""
    row_count, col_count = 3, 2
    captured_table = svc._build_table_grid(
        {
            "rowCount": row_count,
            "colCount": col_count,
            "cells": [
                {"rowIndex": 0, "colIndex": 0, "rowSpan": 2, "colSpan": 2, "text": "block"},
                # Distinct anchor from "block", but its footprint (1,0)-(1,1)
                # collides with the region "block" already merged.
                {"rowIndex": 1, "colIndex": 0, "rowSpan": 1, "colSpan": 2, "text": "conflict"},
                {"rowIndex": 2, "colIndex": 0, "rowSpan": 1, "colSpan": 1, "text": "R2C0"},
                {"rowIndex": 2, "colIndex": 1, "rowSpan": 1, "colSpan": 1, "text": "R2C1"},
            ],
        }
    )
    assert captured_table is not None

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    graphic_frame = slide.shapes.add_table(
        row_count, col_count, Emu(0), Emu(0), Emu(4000000), Emu(2000000)
    )
    table = graphic_frame.table

    svc._apply_table_content(table, captured_table)
    # Must not raise despite the second merge colliding with the first.
    svc._apply_table_merges(table, captured_table)

    # The first merge succeeds (python-pptx concatenates the consumed cell's
    # pre-written text into the origin rather than discarding it - confirmed
    # directly against the real API); the second, colliding merge is the one
    # that must be skipped without raising. Cells outside the conflict are
    # untouched either way.
    assert table.cell(0, 0).text == "block\nconflict"
    assert table.cell(2, 0).text == "R2C0"
    assert table.cell(2, 1).text == "R2C1"


# ---------------------------------------------------------------------------
# Styling: fill, font, header/banding, column widths, row heights
# ---------------------------------------------------------------------------


def test_styling_applies_fill_font_color_size_bold_and_leaves_banding_off():
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(0, 0, 400, 300)
    captured_table = {
        "boundingBox": {"left": 0, "top": 0, "width": 400, "height": 300},
        "rowCount": 2,
        "colCount": 2,
        "cells": [
            {
                "rowIndex": 0,
                "colIndex": 0,
                "rowSpan": 1,
                "colSpan": 1,
                "text": "Header",
                "isHeaderCell": True,
                "backgroundColor": "#3366FF",
                "textColor": "#FFFFFF",
                "fontFamily": "Arial",
                "fontSize": 14,
                "fontWeight": 700,
                "textAlign": "center",
            },
            {"rowIndex": 0, "colIndex": 1, "rowSpan": 1, "colSpan": 1, "text": "H2"},
            {"rowIndex": 1, "colIndex": 0, "rowSpan": 1, "colSpan": 1, "text": "A"},
            {"rowIndex": 1, "colIndex": 1, "rowSpan": 1, "colSpan": 1, "text": "B"},
        ],
    }
    claimed: set = set()
    upgraded = svc._try_upgrade_one_table(slide, captured_table, emu_x, emu_y, claimed)
    assert upgraded is True

    table = slide.shapes[0].table
    assert table.first_row is True
    # Deliberately off: real cell fills (set below) always win over table-
    # style banding in OOXML, so this flag never affects a table with real
    # captured colors either way - and leaving it on would fabricate a
    # stripe pattern for any OTHER table that has no real per-row colors of
    # its own, which is exactly the bug a real user reported.
    assert table.horz_banding is False

    header_cell = table.cell(0, 0)
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    run = header_cell.text_frame.paragraphs[0].runs[0]
    assert run.font.color.rgb == RGBColor.from_string("FFFFFF")
    assert run.font.bold is True
    assert run.font.name == "Arial"
    assert run.font.size.pt == pytest.approx(10.5)
    assert header_cell.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER


def test_styling_applies_normalized_column_widths_and_row_heights():
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(0, 0, 400, 300)
    captured_table = _simple_table(2, 3)
    captured_table["boundingBox"] = {"left": 0, "top": 0, "width": 400, "height": 300}
    captured_table["columnWidthsPx"] = [40, 120, 240]  # 40+120+240=400
    captured_table["rowHeightsPx"] = [100, 200]  # sums to 300
    claimed: set = set()
    upgraded = svc._try_upgrade_one_table(slide, captured_table, emu_x, emu_y, claimed)
    assert upgraded is True

    table = slide.shapes[0].table
    total_width = sum(c.width for c in table.columns)
    total_height = sum(r.height for r in table.rows)
    # Widths/heights must sum exactly to the frame's real EMU size - the
    # whole point of largest-remainder rounding - and be roughly
    # proportional to the captured px ratios (40:120:240 = 1:3:6).
    assert sum(table.columns[i].width for i in range(3)) == total_width
    assert sum(table.rows[i].height for i in range(2)) == total_height
    ratio_0_to_2 = table.columns[0].width / table.columns[2].width
    assert ratio_0_to_2 == pytest.approx(40 / 240, rel=0.05)


def test_styling_falls_back_to_equal_widths_when_every_row_has_a_colspan():
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(0, 0, 400, 300)
    captured_table = {
        "boundingBox": {"left": 0, "top": 0, "width": 400, "height": 300},
        "rowCount": 2,
        "colCount": 2,
        "cells": [
            {"rowIndex": 0, "colIndex": 0, "rowSpan": 1, "colSpan": 2, "text": "wide"},
            {"rowIndex": 1, "colIndex": 0, "rowSpan": 1, "colSpan": 2, "text": "wide2"},
        ],
        # No usable columnWidthsPx since colCount=2 but this is a degenerate
        # case anyway (styling code checks len(columnWidthsPx) == colCount).
        "columnWidthsPx": [200],
    }
    claimed: set = set()
    upgraded = svc._try_upgrade_one_table(slide, captured_table, emu_x, emu_y, claimed)
    assert upgraded is True
    # Must not raise, and the table's default equal-width columns stand.
    table = slide.shapes[0].table
    assert table.columns[0].width == table.columns[1].width


def test_empty_cell_text_skips_run_but_still_applies_fill_and_alignment():
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(0, 0, 400, 300)
    captured_table = {
        "boundingBox": {"left": 0, "top": 0, "width": 400, "height": 300},
        "rowCount": 1,
        "colCount": 1,
        "cells": [
            {
                "rowIndex": 0,
                "colIndex": 0,
                "rowSpan": 1,
                "colSpan": 1,
                "text": "",
                "backgroundColor": "#FF0000",
                "textAlign": "right",
            }
        ],
    }
    claimed: set = set()
    upgraded = svc._try_upgrade_one_table(slide, captured_table, emu_x, emu_y, claimed)
    assert upgraded is True

    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    table = slide.shapes[0].table
    cell = table.cell(0, 0)
    assert len(cell.text_frame.paragraphs[0].runs) == 0
    assert cell.fill.fore_color.rgb == RGBColor.from_string("FF0000")
    assert cell.text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT


# ---------------------------------------------------------------------------
# Row-separator bottom border - a real user request (identical to the
# original app's `border-b border-black/15` row lines), built by hand-
# authoring <a:lnB> since python-pptx has zero API for cell borders
# (confirmed directly against pptx.table._Cell/Table's full public
# attribute lists and pptx.oxml.table.CT_TableCellProperties's own schema
# declaration - no line elements modeled anywhere in this library).
# ---------------------------------------------------------------------------

_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def test_apply_cell_bottom_border_emits_lnb_before_existing_fill():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    gf = slide.shapes.add_table(1, 1, Emu(0), Emu(0), Emu(1000000), Emu(500000))
    cell = gf.table.cell(0, 0)
    cell.fill.solid()
    from pptx.dml.color import RGBColor

    cell.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")

    svc._apply_cell_bottom_border(cell, "000000", 0.15, 12700)

    tcPr = cell._tc.find(f"{_A_NS}tcPr")
    children = list(tcPr)
    tags = [etree.QName(c).localname for c in children]
    assert tags[0] == "lnB"
    assert tags[1] == "solidFill"

    lnB = children[0]
    assert lnB.get("w") == "12700"
    srgbClr = lnB.find(f"{_A_NS}solidFill/{_A_NS}srgbClr")
    assert srgbClr.get("val") == "000000"
    alpha = srgbClr.find(f"{_A_NS}alpha")
    assert alpha.get("val") == "15000"


def test_apply_cell_bottom_border_with_no_alpha_emits_no_alpha_element():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    gf = slide.shapes.add_table(1, 1, Emu(0), Emu(0), Emu(1000000), Emu(500000))
    cell = gf.table.cell(0, 0)

    svc._apply_cell_bottom_border(cell, "FF0000", None, 9525)

    tcPr = cell._tc.find(f"{_A_NS}tcPr")
    srgbClr = tcPr.find(f"{_A_NS}lnB/{_A_NS}solidFill/{_A_NS}srgbClr")
    assert srgbClr.get("val") == "FF0000"
    assert srgbClr.find(f"{_A_NS}alpha") is None


def test_apply_cell_bottom_border_replaces_existing_lnb_without_duplicating():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    gf = slide.shapes.add_table(1, 1, Emu(0), Emu(0), Emu(1000000), Emu(500000))
    cell = gf.table.cell(0, 0)

    svc._apply_cell_bottom_border(cell, "000000", None, 12700)
    svc._apply_cell_bottom_border(cell, "FFFFFF", None, 25400)

    tcPr = cell._tc.find(f"{_A_NS}tcPr")
    lnBs = tcPr.findall(f"{_A_NS}lnB")
    assert len(lnBs) == 1
    assert lnBs[0].get("w") == "25400"


def test_apply_cell_bottom_border_survives_save_and_reload():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    gf = slide.shapes.add_table(1, 1, Emu(0), Emu(0), Emu(1000000), Emu(500000))
    cell = gf.table.cell(0, 0)
    cell.text = "Row"
    svc._apply_cell_bottom_border(cell, "000000", 0.15, 12700)

    import tempfile
    import os

    fd, path = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    try:
        prs.save(path)
        reopened = Presentation(path)
        rcell = reopened.slides[0].shapes[0].table.cell(0, 0)
        assert rcell.text == "Row"
        tcPr = rcell._tc.find(f"{_A_NS}tcPr")
        assert tcPr.find(f"{_A_NS}lnB") is not None
    finally:
        os.remove(path)


def test_styling_applies_captured_bottom_border_end_to_end():
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(0, 0, 400, 300)
    captured_table = {
        "boundingBox": {"left": 0, "top": 0, "width": 400, "height": 300},
        "rowCount": 2,
        "colCount": 2,
        "cells": [
            {
                "rowIndex": 0,
                "colIndex": 0,
                "rowSpan": 1,
                "colSpan": 1,
                "text": "A",
                "bottomBorderColor": "rgba(0, 0, 0, 0.15)",
                "bottomBorderWidthPx": 1,
            },
            {"rowIndex": 0, "colIndex": 1, "rowSpan": 1, "colSpan": 1, "text": "B"},
            {"rowIndex": 1, "colIndex": 0, "rowSpan": 1, "colSpan": 1, "text": "C"},
            {"rowIndex": 1, "colIndex": 1, "rowSpan": 1, "colSpan": 1, "text": "D"},
        ],
    }
    claimed: set = set()
    upgraded = svc._try_upgrade_one_table(slide, captured_table, emu_x, emu_y, claimed)
    assert upgraded is True

    table = slide.shapes[0].table
    bordered_cell = table.cell(0, 0)._tc.find(f"{_A_NS}tcPr/{_A_NS}lnB")
    assert bordered_cell is not None
    unbordered_cell = table.cell(0, 1)
    tcPr = unbordered_cell._tc.find(f"{_A_NS}tcPr")
    assert tcPr is None or tcPr.find(f"{_A_NS}lnB") is None


@pytest.mark.parametrize("bad_width", [0, -1, 0.1, 21, "2", None, True])
def test_styling_ignores_out_of_range_or_malformed_border_width(bad_width):
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(0, 0, 400, 300)
    captured_table = {
        "boundingBox": {"left": 0, "top": 0, "width": 400, "height": 300},
        "rowCount": 1,
        "colCount": 1,
        "cells": [
            {
                "rowIndex": 0,
                "colIndex": 0,
                "rowSpan": 1,
                "colSpan": 1,
                "text": "A",
                "bottomBorderColor": "#000000",
                "bottomBorderWidthPx": bad_width,
            }
        ],
    }
    claimed: set = set()
    upgraded = svc._try_upgrade_one_table(slide, captured_table, emu_x, emu_y, claimed)
    assert upgraded is True
    # Must not raise, and must not apply a border for a bad width value.
    table = slide.shapes[0].table
    tcPr = table.cell(0, 0)._tc.find(f"{_A_NS}tcPr")
    assert tcPr is None or tcPr.find(f"{_A_NS}lnB") is None


# ---------------------------------------------------------------------------
# upgrade_flattened_tables_to_native - end-to-end
# ---------------------------------------------------------------------------


def test_upgrade_flattened_tables_to_native_noop_without_token():
    asyncio.run(
        svc.upgrade_flattened_tables_to_native("/nonexistent/path.pptx", None, uuid.uuid4())
    )


def test_upgrade_flattened_tables_to_native_noop_when_nothing_captured(monkeypatch):
    monkeypatch.setattr(svc.table_capture_store, "take_capture", lambda token: None)
    asyncio.run(
        svc.upgrade_flattened_tables_to_native(
            "/nonexistent/path.pptx", "some-token", uuid.uuid4()
        )
    )


def test_upgrade_flattened_tables_to_native_end_to_end(monkeypatch, tmp_path):
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(100, 100, 400, 300)
    pptx_path = tmp_path / "deck.pptx"
    prs.save(str(pptx_path))

    captured = {
        "presentation_id": "irrelevant",
        "tables": [
            {
                "slideOrderIndex": 0,
                "boundingBox": {"left": 100, "top": 100, "width": 400, "height": 300},
                "rowCount": 2,
                "colCount": 2,
                "cells": [
                    {"rowIndex": 0, "colIndex": 0, "rowSpan": 1, "colSpan": 1, "text": "A"},
                    {"rowIndex": 0, "colIndex": 1, "rowSpan": 1, "colSpan": 1, "text": "B"},
                    {"rowIndex": 1, "colIndex": 0, "rowSpan": 1, "colSpan": 1, "text": "C"},
                    {"rowIndex": 1, "colIndex": 1, "rowSpan": 1, "colSpan": 1, "text": "D"},
                ],
            }
        ],
    }
    monkeypatch.setattr(svc.table_capture_store, "take_capture", lambda token: captured)

    async def _fake_expected_count(presentation_id):
        return 1

    monkeypatch.setattr(svc, "_expected_slide_count", _fake_expected_count)

    asyncio.run(
        svc.upgrade_flattened_tables_to_native(str(pptx_path), "tok", uuid.uuid4())
    )

    reopened = Presentation(str(pptx_path))
    shapes = list(reopened.slides[0].shapes)
    assert shapes[0].shape_type == MSO_SHAPE_TYPE.TABLE
    assert shapes[0].table.cell(0, 0).text == "A"


def test_upgrade_flattened_tables_to_native_slide_count_mismatch_is_untouched(
    monkeypatch, tmp_path
):
    prs, slide, picture, emu_x, emu_y = _make_presentation_with_picture(100, 100, 400, 300)
    pptx_path = tmp_path / "deck.pptx"
    prs.save(str(pptx_path))
    original_bytes = pptx_path.read_bytes()

    captured = {
        "presentation_id": "irrelevant",
        "tables": [
            {
                "slideOrderIndex": 0,
                "boundingBox": {"left": 100, "top": 100, "width": 400, "height": 300},
                "rowCount": 1,
                "colCount": 1,
                "cells": [{"rowIndex": 0, "colIndex": 0, "rowSpan": 1, "colSpan": 1, "text": "A"}],
            }
        ],
    }
    monkeypatch.setattr(svc.table_capture_store, "take_capture", lambda token: captured)

    async def _fake_expected_count(presentation_id):
        return 2  # deliberately wrong -> should bail out entirely

    monkeypatch.setattr(svc, "_expected_slide_count", _fake_expected_count)

    asyncio.run(svc.upgrade_flattened_tables_to_native(str(pptx_path), "tok", uuid.uuid4()))

    assert pptx_path.read_bytes() == original_bytes


def test_two_pictures_on_same_slide_each_claimed_independently_by_a_table_and_a_chart():
    """A slide carrying both a table and a chart (the stress-test deck's
    "Combination slide" pattern) must let each pass claim its own picture
    without double-claiming - per-slide claimed_shape_ids sets, allocated
    fresh for each pass, are what make this safe (see Design Decision 3)."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    emu_x = prs.slide_width / 1280
    emu_y = prs.slide_height / 720

    def _add_pic(left, top, width, height):
        img = Image.new("RGB", (10, 10), color="green")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        buf.seek(0)
        return slide.shapes.add_picture(
            buf,
            Emu(int(left * emu_x)),
            Emu(int(top * emu_y)),
            Emu(int(width * emu_x)),
            Emu(int(height * emu_y)),
        )

    table_picture = _add_pic(0, 0, 300, 200)
    chart_picture_stand_in = _add_pic(700, 0, 300, 200)

    captured_table = _simple_table(2, 2)
    captured_table["boundingBox"] = {"left": 0, "top": 0, "width": 300, "height": 200}

    claimed: set = set()
    upgraded = svc._try_upgrade_one_table(slide, captured_table, emu_x, emu_y, claimed)
    assert upgraded is True

    shapes = list(slide.shapes)
    types = [s.shape_type for s in shapes]
    assert types.count(MSO_SHAPE_TYPE.TABLE) == 1
    assert types.count(MSO_SHAPE_TYPE.PICTURE) == 1  # the other picture untouched
