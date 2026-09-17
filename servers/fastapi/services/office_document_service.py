import io
import logging
import os
import posixpath
import re
import zipfile
from pathlib import Path
from typing import List, Optional
from xml.etree import ElementTree

import openpyxl
from openpyxl.utils import range_boundaries

from models.pptx_chart_data import ExtractedChart, ExtractedChartSeries, PptxStructuredData
from templates.v2.models.elements import ChartType

LOGGER = logging.getLogger(__name__)


class OfficeDocumentError(Exception):
    pass


_DOCX_EXTENSIONS = {".docx", ".docm"}
_PPTX_EXTENSIONS = {".pptx", ".pptm"}
_XLSX_EXTENSIONS = {".xlsx", ".xlsm"}
_ODF_EXTENSIONS = {".odt", ".odp", ".ods"}
_TEXT_EXTENSIONS = {".csv", ".tsv"}
_UNSUPPORTED_LEGACY_EXTENSIONS = {".doc", ".ppt", ".xls", ".rtf"}


def _natural_key(value: str) -> list[object]:
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", value)
    ]


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _read_xml(archive: zipfile.ZipFile, member: str) -> ElementTree.Element:
    try:
        return ElementTree.fromstring(archive.read(member))
    except (KeyError, ElementTree.ParseError) as exc:
        raise OfficeDocumentError(f"Could not read {member}") from exc


def _text_nodes(root: ElementTree.Element) -> list[str]:
    return [
        value
        for element in root.iter()
        if _local_name(element.tag) == "t"
        and (value := (element.text or "").strip())
    ]


def _extract_docx(archive: zipfile.ZipFile) -> str:
    root = _read_xml(archive, "word/document.xml")
    paragraphs: list[str] = []
    for paragraph in root.iter():
        if _local_name(paragraph.tag) != "p":
            continue
        text = " ".join(_text_nodes(paragraph))
        if text:
            paragraphs.append(text)
    return "\n".join(paragraphs)


def _extract_pptx(archive: zipfile.ZipFile) -> str:
    slide_members = sorted(
        (
            name
            for name in archive.namelist()
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        ),
        key=_natural_key,
    )
    slides = [" ".join(_text_nodes(_read_xml(archive, member))) for member in slide_members]
    return "\n\n".join(slide for slide in slides if slide)


def _find_children(element: ElementTree.Element, local_name: str) -> list[ElementTree.Element]:
    return [child for child in element if _local_name(child.tag) == local_name]


def _find_first(
    element: Optional[ElementTree.Element], local_name: str
) -> Optional[ElementTree.Element]:
    if element is None:
        return None
    children = _find_children(element, local_name)
    return children[0] if children else None


def _find_descendant(
    element: Optional[ElementTree.Element], local_name: str
) -> Optional[ElementTree.Element]:
    if element is None:
        return None
    for node in element.iter():
        if _local_name(node.tag) == local_name:
            return node
    return None


def _read_rels(archive: zipfile.ZipFile, rels_member: str) -> list[dict]:
    if rels_member not in archive.namelist():
        return []
    root = _read_xml(archive, rels_member)
    return [child.attrib for child in root if _local_name(child.tag) == "Relationship"]


def _resolve_relative_target(base_dir: str, target: str) -> str:
    return posixpath.normpath(posixpath.join(base_dir, target))


def _map_charts_to_slides(archive: zipfile.ZipFile) -> dict[str, int]:
    mapping: dict[str, int] = {}
    slide_rels_members = sorted(
        (
            name
            for name in archive.namelist()
            if re.fullmatch(r"ppt/slides/_rels/slide\d+\.xml\.rels", name)
        ),
        key=_natural_key,
    )
    for rels_member in slide_rels_members:
        match = re.search(r"slide(\d+)\.xml\.rels$", rels_member)
        if not match:
            continue
        slide_index = int(match.group(1))
        for rel in _read_rels(archive, rels_member):
            target = rel.get("Target", "")
            if "charts/chart" in target and target.endswith(".xml"):
                chart_member = _resolve_relative_target("ppt/slides", target)
                mapping[chart_member] = slide_index
    return mapping


def _find_embedded_workbook_path(archive: zipfile.ZipFile, chart_member: str) -> Optional[str]:
    base_dir = posixpath.dirname(chart_member)
    rels_member = posixpath.join(base_dir, "_rels", posixpath.basename(chart_member) + ".rels")
    for rel in _read_rels(archive, rels_member):
        rel_type = rel.get("Type", "")
        target = rel.get("Target", "")
        if "package" in rel_type.lower() or target.lower().endswith((".xlsx", ".xlsm")):
            return _resolve_relative_target(base_dir, target)
    return None


def _pt_pairs(container: ElementTree.Element) -> list[tuple[int, Optional[str]]]:
    pairs: list[tuple[int, Optional[str]]] = []
    for pt in _find_children(container, "pt"):
        idx_raw = pt.attrib.get("idx")
        if idx_raw is None or not idx_raw.isdigit():
            continue
        v_el = _find_first(pt, "v")
        value = (v_el.text or "").strip() if v_el is not None and v_el.text else None
        pairs.append((int(idx_raw), value))
    return pairs


def _cache_pairs(ref_element: Optional[ElementTree.Element]) -> list[tuple[int, Optional[str]]]:
    """Ordered (idx, value) pairs from the strCache/numCache under a strRef/numRef."""
    if ref_element is None:
        return []
    cache = _find_descendant(ref_element, "strCache") or _find_descendant(ref_element, "numCache")
    if cache is None:
        return []
    return _pt_pairs(cache)


def _multi_lvl_count(multi_lvl_ref: Optional[ElementTree.Element]) -> int:
    cache = _find_first(multi_lvl_ref, "multiLvlStrCache")
    pt_count_el = _find_first(cache, "ptCount")
    if pt_count_el is not None and pt_count_el.attrib.get("val", "").isdigit():
        return int(pt_count_el.attrib["val"])
    return 0


def _multi_level_category_labels(
    multi_lvl_ref: Optional[ElementTree.Element], count: int
) -> list[Optional[str]]:
    """Flatten a multiLvlStrRef's category levels into one "outer / inner" label per index.

    Levels are ordered innermost-first in the XML (ECMA-376 CT_MultiLvlStrRef); a group label
    only appears at its first index and implicitly applies to the following indices until the
    next label, so each level is forward-filled before combining.
    """
    cache = _find_first(multi_lvl_ref, "multiLvlStrCache")
    if cache is None:
        return [None] * count
    levels = _find_children(cache, "lvl")
    if not levels:
        return [None] * count

    filled_levels: list[list[Optional[str]]] = []
    for level in levels:
        pairs = dict(_pt_pairs(level))
        filled: list[Optional[str]] = []
        last: Optional[str] = None
        for i in range(count):
            if pairs.get(i):
                last = pairs[i]
            filled.append(last)
        filled_levels.append(filled)

    labels: list[Optional[str]] = []
    for i in range(count):
        parts = [level[i] for level in reversed(filled_levels) if level[i]]
        labels.append(" / ".join(parts) if parts else None)
    return labels


def _cache_count(ref_element: Optional[ElementTree.Element]) -> int:
    if ref_element is None:
        return 0
    cache = _find_descendant(ref_element, "strCache") or _find_descendant(ref_element, "numCache")
    pt_count_el = _find_first(cache, "ptCount") if cache is not None else None
    if pt_count_el is not None and pt_count_el.attrib.get("val", "").isdigit():
        return int(pt_count_el.attrib["val"])
    pairs = _cache_pairs(ref_element)
    return (max(idx for idx, _ in pairs) + 1) if pairs else 0


def _ordered_values(ref_element: Optional[ElementTree.Element], count: int) -> list[Optional[str]]:
    pairs = dict(_cache_pairs(ref_element))
    return [pairs.get(i) for i in range(count)]


def _formula_text(ref_element: Optional[ElementTree.Element]) -> Optional[str]:
    f_el = _find_first(ref_element, "f")
    return f_el.text.strip() if f_el is not None and f_el.text else None


def _series_name(ser_element: ElementTree.Element, index: int) -> str:
    default = f"Series {index + 1}"
    tx_el = _find_first(ser_element, "tx")
    if tx_el is None:
        return default
    str_ref = _find_first(tx_el, "strRef")
    if str_ref is not None:
        pairs = _cache_pairs(str_ref)
        if pairs and pairs[0][1]:
            return pairs[0][1]
    v_el = _find_first(tx_el, "v")
    if v_el is not None and v_el.text:
        return v_el.text.strip()
    return default


_BAR_CHART_LOCAL_NAMES = {"barChart", "bar3DChart"}
_CHART_TYPE_MAP = {
    "lineChart": ChartType.LINE,
    "line3DChart": ChartType.LINE,
    "pieChart": ChartType.PIE,
    "pie3DChart": ChartType.PIE,
    "doughnutChart": ChartType.DONUT,
    "radarChart": ChartType.RADAR,
    "areaChart": ChartType.AREA,
    "area3DChart": ChartType.AREA,
    "scatterChart": ChartType.SCATTER,
}


def _detect_chart_type_element(
    plot_area: ElementTree.Element,
) -> tuple[Optional[ElementTree.Element], str, Optional[ChartType]]:
    for child in plot_area:
        local = _local_name(child.tag)
        if not local.endswith("Chart"):
            continue
        if local in _BAR_CHART_LOCAL_NAMES:
            bar_dir_el = _find_descendant(child, "barDir")
            grouping_el = _find_descendant(child, "grouping")
            bar_dir = bar_dir_el.attrib.get("val", "col") if bar_dir_el is not None else "col"
            grouping = (
                grouping_el.attrib.get("val", "clustered")
                if grouping_el is not None
                else "clustered"
            )
            is_horizontal = bar_dir == "bar"
            is_stacked = grouping in {"stacked", "percentStacked"}
            raw_type = f"{local}/{bar_dir}/{grouping}"
            if is_horizontal and is_stacked:
                mapped = ChartType.HORIZONTAL_STACKED_BAR
            elif is_horizontal:
                mapped = ChartType.HORIZONTAL_BAR
            elif is_stacked:
                mapped = ChartType.STACKED_BAR
            else:
                mapped = ChartType.BAR
            return child, raw_type, mapped
        return child, local, _CHART_TYPE_MAP.get(local)
    return None, "unknown", None


def _chart_title(chart_element: Optional[ElementTree.Element]) -> Optional[str]:
    title_el = _find_first(chart_element, "title")
    if title_el is None:
        return None
    joined = " ".join(_text_nodes(title_el)).strip()
    return joined or None


def _parse_formula_ref(formula: str) -> tuple[Optional[str], Optional[str]]:
    if not formula or "!" not in formula:
        return None, None
    sheet_part, range_part = formula.rsplit("!", 1)
    sheet_name = sheet_part.strip("'")
    cell_range = range_part.replace("$", "")
    return sheet_name, cell_range


def _resolve_workbook_range(workbook, sheet_name: Optional[str], cell_range: str) -> Optional[list]:
    try:
        worksheet = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook.active
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    except (KeyError, ValueError):
        return None
    values = []
    if max_row > min_row and max_col == min_col:
        for row in range(min_row, max_row + 1):
            values.append(worksheet.cell(row=row, column=min_col).value)
    elif max_col > min_col and max_row == min_row:
        for col in range(min_col, max_col + 1):
            values.append(worksheet.cell(row=min_row, column=col).value)
    else:
        values.append(worksheet.cell(row=min_row, column=min_col).value)
    return values


def _try_workbook_values(
    archive: zipfile.ZipFile, chart_member: str, formula: Optional[str]
) -> Optional[list]:
    if not formula:
        return None
    workbook_path = _find_embedded_workbook_path(archive, chart_member)
    if workbook_path is None or workbook_path not in archive.namelist():
        return None
    sheet_name, cell_range = _parse_formula_ref(formula)
    if cell_range is None:
        return None
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(archive.read(workbook_path)), data_only=True, read_only=True
        )
    except Exception:
        return None
    try:
        return _resolve_workbook_range(workbook, sheet_name, cell_range)
    except Exception:
        return None
    finally:
        workbook.close()


def _extract_single_chart(
    archive: zipfile.ZipFile,
    chart_member: str,
    source_file: str,
    slide_index: Optional[int],
) -> ExtractedChart:
    root = _read_xml(archive, chart_member)
    chart_el = _find_descendant(root, "chart")
    plot_area = _find_descendant(chart_el, "plotArea")
    title = _chart_title(chart_el)

    if plot_area is None:
        return ExtractedChart(
            source_file=source_file,
            slide_index=slide_index,
            chart_part=chart_member,
            title=title,
            raw_ooxml_type="unknown",
            extraction_method="none",
            extractable=False,
        )

    chart_type_el, raw_type, chart_type = _detect_chart_type_element(plot_area)
    if chart_type_el is None:
        return ExtractedChart(
            source_file=source_file,
            slide_index=slide_index,
            chart_part=chart_member,
            title=title,
            raw_ooxml_type=raw_type,
            extraction_method="none",
            extractable=False,
        )

    categories: list[str] = []
    series_list: list[ExtractedChartSeries] = []
    used_workbook = False

    parsed_series: list[dict] = []
    for index, ser in enumerate(_find_children(chart_type_el, "ser")):
        cat_el = _find_first(ser, "cat")
        val_el = _find_first(ser, "val")
        cat_ref = (
            (_find_first(cat_el, "strRef") or _find_first(cat_el, "numRef"))
            if cat_el is not None
            else None
        )
        multi_lvl_cat_ref = (
            _find_first(cat_el, "multiLvlStrRef") if cat_el is not None and cat_ref is None else None
        )
        val_ref = _find_first(val_el, "numRef") if val_el is not None else None
        if val_ref is None:
            continue

        count = max(
            _cache_count(cat_ref),
            _multi_lvl_count(multi_lvl_cat_ref),
            _cache_count(val_ref),
        )
        parsed_series.append(
            {
                "ser": ser,
                "index": index,
                "cat_ref": cat_ref,
                "multi_lvl_cat_ref": multi_lvl_cat_ref,
                "val_ref": val_ref,
                "count": count,
            }
        )

    any_series = bool(parsed_series)

    if parsed_series:
        # Categories are shared across every series on a chart (ECMA-376), so they must be
        # extracted using the widest point count seen across ALL series' caches, not just the
        # series that happens to populate `categories` first — otherwise a series with a larger
        # (possibly less stale) cached count than that first series ends up with a values list
        # longer than `categories`, silently misaligning category/value pairs downstream
        # (document_fact_dedup_service.py's `zip(chart.categories, series.values)`).
        chart_count = max(item["count"] for item in parsed_series)

        first = parsed_series[0]
        if first["multi_lvl_cat_ref"] is not None:
            raw_categories: list = list(
                _multi_level_category_labels(first["multi_lvl_cat_ref"], chart_count)
            )
        elif first["cat_ref"] is not None:
            raw_categories = list(_ordered_values(first["cat_ref"], chart_count))
            workbook_categories = _try_workbook_values(
                archive, chart_member, _formula_text(first["cat_ref"])
            )
            if (
                workbook_categories is not None
                and len(workbook_categories) == chart_count
                and chart_count > 0
            ):
                raw_categories = workbook_categories
                used_workbook = True
        else:
            raw_categories = [None] * chart_count
        categories = ["" if value is None else str(value) for value in raw_categories]

        for item in parsed_series:
            val_ref = item["val_ref"]
            raw_values: list = list(_ordered_values(val_ref, chart_count))
            workbook_values = _try_workbook_values(archive, chart_member, _formula_text(val_ref))
            if workbook_values is not None and len(workbook_values) == chart_count and chart_count > 0:
                raw_values = workbook_values
                used_workbook = True

            numeric_values: list[Optional[float]] = []
            for value in raw_values:
                if value is None or value == "":
                    numeric_values.append(None)
                    continue
                try:
                    numeric_values.append(float(value))
                except (TypeError, ValueError):
                    numeric_values.append(None)

            series_list.append(
                ExtractedChartSeries(name=_series_name(item["ser"], item["index"]), values=numeric_values)
            )

    extractable = any_series and any(
        value is not None for series in series_list for value in series.values
    )
    extraction_method = "none"
    if extractable:
        extraction_method = "embedded_workbook" if used_workbook else "chart_xml_cache"

    return ExtractedChart(
        source_file=source_file,
        slide_index=slide_index,
        chart_part=chart_member,
        title=title,
        chart_type=chart_type,
        raw_ooxml_type=raw_type,
        categories=categories,
        series=series_list,
        extraction_method=extraction_method,
        extractable=extractable,
    )


def extract_pptx_structured_data(file_path: str) -> PptxStructuredData:
    """Pull real, editable chart data out of a pptx's native OOXML charts.

    `_extract_pptx` above only reads slide text runs and never touches `ppt/charts/*.xml` or
    the embedded workbooks, so a native chart's real series data is otherwise lost entirely.
    """
    source_file = os.path.basename(file_path)
    charts: list[ExtractedChart] = []
    try:
        with zipfile.ZipFile(file_path) as archive:
            chart_members = sorted(
                (
                    name
                    for name in archive.namelist()
                    if re.fullmatch(r"ppt/charts/chart\d+\.xml", name)
                ),
                key=_natural_key,
            )
            slide_index_by_chart = _map_charts_to_slides(archive)
            for chart_member in chart_members:
                try:
                    charts.append(
                        _extract_single_chart(
                            archive,
                            chart_member,
                            source_file,
                            slide_index_by_chart.get(chart_member),
                        )
                    )
                except Exception:
                    LOGGER.warning(
                        "Failed to extract chart data from %s in %s",
                        chart_member,
                        file_path,
                        exc_info=True,
                    )
                    charts.append(
                        ExtractedChart(
                            source_file=source_file,
                            slide_index=slide_index_by_chart.get(chart_member),
                            chart_part=chart_member,
                            raw_ooxml_type="unknown",
                            extraction_method="none",
                            extractable=False,
                        )
                    )
    except (OSError, zipfile.BadZipFile):
        return PptxStructuredData(charts=[])
    return PptxStructuredData(charts=charts)


def _extract_shared_strings(archive: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in archive.namelist():
        return []
    root = _read_xml(archive, "xl/sharedStrings.xml")
    return [
        " ".join(_text_nodes(item))
        for item in root.iter()
        if _local_name(item.tag) == "si"
    ]


def _extract_xlsx(archive: zipfile.ZipFile) -> str:
    shared_strings = _extract_shared_strings(archive)
    sheet_members = sorted(
        (
            name
            for name in archive.namelist()
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name)
        ),
        key=_natural_key,
    )
    sheets: list[str] = []
    for member in sheet_members:
        root = _read_xml(archive, member)
        rows: list[str] = []
        for row in root.iter():
            if _local_name(row.tag) != "row":
                continue
            values: list[str] = []
            for cell in row:
                if _local_name(cell.tag) != "c":
                    continue
                cell_type = cell.attrib.get("t")
                if cell_type == "inlineStr":
                    value = " ".join(_text_nodes(cell))
                else:
                    raw_value = next(
                        (
                            (element.text or "").strip()
                            for element in cell
                            if _local_name(element.tag) == "v"
                        ),
                        "",
                    )
                    if cell_type == "s" and raw_value.isdigit():
                        index = int(raw_value)
                        value = (
                            shared_strings[index]
                            if index < len(shared_strings)
                            else raw_value
                        )
                    else:
                        value = raw_value
                if value:
                    values.append(value)
            if values:
                rows.append("\t".join(values))
        if rows:
            sheets.append("\n".join(rows))
    return "\n\n".join(sheets)


def _extract_odf(archive: zipfile.ZipFile) -> str:
    root = _read_xml(archive, "content.xml")
    values = [
        " ".join(text.strip() for text in element.itertext() if text.strip())
        for element in root.iter()
        if _local_name(element.tag) in {"p", "h"}
    ]
    return "\n".join(value for value in values if value)


def extract_office_document_text(file_path: str) -> str:
    extension = Path(file_path).suffix.lower()
    if extension in _TEXT_EXTENSIONS:
        with open(file_path, "r", encoding="utf-8", errors="replace") as file:
            return file.read()

    if extension in _UNSUPPORTED_LEGACY_EXTENSIONS:
        raise OfficeDocumentError(
            f"{extension} files require an external office conversion engine; "
            "save the document in a modern OOXML or OpenDocument format first"
        )

    try:
        with zipfile.ZipFile(file_path) as archive:
            if extension in _DOCX_EXTENSIONS:
                return _extract_docx(archive)
            if extension in _PPTX_EXTENSIONS:
                return _extract_pptx(archive)
            if extension in _XLSX_EXTENSIONS:
                return _extract_xlsx(archive)
            if extension in _ODF_EXTENSIONS:
                return _extract_odf(archive)
    except (OSError, zipfile.BadZipFile) as exc:
        raise OfficeDocumentError(
            f"Could not parse {os.path.basename(file_path)}"
        ) from exc

    raise OfficeDocumentError(f"Unsupported office document format: {extension}")
